#!/usr/bin/env python3
"""엑셀 렌더링을 **LibreOffice 가 그린 화면 글자**와 대조한다.

  python .claude/skills/doc-archive/scripts/check_against_libreoffice.py [--cases N]

종료코드: 0 어긋남 없음 / 1 어긋난 셀 있음 / 2 LibreOffice 를 못 찾아 **재지 못함**

**왜 있나 — 「감사할 수 없음」을 없애려고 만들었다.** `xlsx_to_blocks.py` 의 구역 선택
(서식의 `;` 로 갈린 양수·음수·0·문자 구역)은 2026-08-22 에 일회용 오라클로 검증했는데
그 오라클을 커밋하지 않아, 감사자가 그 축을 재현할 수 없었다(원장의 PARKED 항목).
오라클은 이미 사라졌으므로 **재현 가능한 쪽**으로 갈아탄다 — 이 스크립트가 그것이다.

**시험(`test_xlsx_to_blocks.py`)과 역할이 다르다.** 시험은 우리가 정한 기대값과 맞는지
보고 의존성이 없다 — 그래서 늘 돌린다. 이것은 **바깥 렌더러**와 맞는지 보는 것이라
LibreOffice 가 필요하고, 값을 우리가 정하지 않는다. 서식 해석을 손볼 때 함께 돌린다.

**LibreOffice 는 엑셀이 아니다.** 대리자일 뿐이라 둘이 갈리는 자리가 있고, 그런 값은
`xlsx_to_blocks` 가 known=False 로 세어 여기서도 「안 그림」으로 빠진다. 실제로 이 대조가
LibreOffice 의 자기모순을 잡아낸 적이 있다 — 68601575893949.84 를 `0.0` 으로는 `…9`,
`0.00` 으로는 `…80` 으로 그렸다(2026-08-23). 그래서 **「안 그림」이 늘어난 것은 고장이
아니다.** 이 검사가 지키는 것은 하나다 — **그린 것 중에 어긋난 것이 없어야 한다.**
"""
import argparse
import csv
import os
import random
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

try:
    import openpyxl
except ImportError:
    # 이 파일은 LibreOffice 가 없을 때 이미 2(못 쟀음)를 낸다. openpyxl 도 같은 갈래다 —
    # 「없어서 못 쟀다」를 「재서 틀렸다」(1)와 가른다.
    print(
        "openpyxl 이 없습니다. pip install -r "
        ".claude/skills/doc-archive/scripts/requirements.txt",
        file=sys.stderr,
    )
    sys.exit(2)

from xlsx_to_blocks import render_value

# 윈도우 파이프로 나가도 한글이 안 깨지게. (`test_xlsx_to_blocks.py` 와 같은 이유)
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

_SOFFICE_CANDIDATES = [
    "soffice",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
]

# CSV 로 내보내면 LibreOffice 는 **화면에 그린 글자**를 쓴다. 그게 우리가 맞춰야 할 대상이다.
#
# **큰따옴표를 넣지 않는다.** 셸에 치던 모양(`csv:"Text - txt - csv (StarCalc)":"44,…"`)을
# 그대로 subprocess 에 넘기면 셸이 벗겨 주던 따옴표가 인자에 남아 soffice 가
# `Please verify input parameters` 로 거절한다 — csv 가 안 생길 뿐 에러 메시지는
# 파일 탓처럼 보인다. 인자 목록으로 넘길 때는 벗긴 모양이 맞다.
#
# 숫자 뜻: 44=쉼표 구분 · 34=따옴표 · 76=UTF-8 · 1=첫 줄부터 · 0=기본 언어
_CSV_FILTER = "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,true"


def find_soffice():
    for c in _SOFFICE_CANDIDATES:
        p = shutil.which(c) if os.sep not in c else (c if Path(c).exists() else None)
        if p:
            return p
    return None


# ── 재는 축 ─────────────────────────────────────────────────────────────
#
# **구역 축이 이 스크립트의 존재 이유다.** 서식은 `;` 로 최대 네 구역(양수·음수·0·문자)을
# 갖고, 값의 부호가 어느 구역을 쓸지 정한다. 2026-08-22 까지 늘 1번 구역만 써서 회계
# 서식의 0 이 `0`(엑셀은 `-`)으로, 괄호 음수가 `-1,520`(엑셀은 `(1,520)`)으로 나갔다.
_SECTION_FORMATS = [
    "#,##0",
    "#,##0;(#,##0)",
    "#,##0;[Red](#,##0)",
    "#,##0;-#,##0;-",
    '_-* #,##0_-;_-* -#,##0_-;_-* "-"_-;_-@_-',   # 한국 회계 서식 (가장 흔하다)
    '_-* #,##0.00_-;_-* -#,##0.00_-;_-* "-"??_-;_-@_-',
    "#,##0;;",                                      # 음수·0 구역이 비었다 (화면은 빈칸)
    "0.0%;[Red](0.0%)",
    "0.00",
    "0.0",
    "0",
    "#,##0.0",
    "#,##0.00",
    "0.0%",
    "0.00%",
]

# 구역을 실제로 갈라 쓰게 하는 값들. 0 과 음수가 빠지면 이 축을 안 재는 것이 된다.
_SECTION_VALUES = [0, 1520.4, -1520.4, 1, -1, 0.125, -0.125, 1234567.89, -1234567.89]


def build_cases(n_random, seed):
    """(값, 서식) 목록. 구역 축은 전수, 반올림 축은 무작위."""
    cases = [(v, f) for f in _SECTION_FORMATS for v in _SECTION_VALUES]

    rnd = random.Random(seed)
    for _ in range(n_random):
        exp = rnd.choice([0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16])
        v = rnd.uniform(10**exp, 10 ** (exp + 1))
        if rnd.random() < 0.3:
            v = -v
        cases.append((v, rnd.choice(_SECTION_FORMATS)))
    return cases


def render_with_libreoffice(soffice, cases, workdir):
    src = workdir / "cases.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, (v, fmt) in enumerate(cases, start=1):
        c = ws.cell(row=i, column=1, value=v)
        c.number_format = fmt
    wb.save(src)

    r = subprocess.run(
        [soffice, "--headless", "--convert-to", _CSV_FILTER, "--outdir", str(workdir), str(src)],
        capture_output=True,
        text=True,
        timeout=300,
    )
    out = workdir / "cases.csv"
    if not out.exists():
        raise RuntimeError(
            f"LibreOffice 가 csv 를 만들지 못했습니다: {(r.stderr or r.stdout).strip()[:300]}"
        )
    with out.open(encoding="utf-8", newline="") as fh:
        drawn = [row[0] if row else "" for row in csv.reader(fh)]

    # **저장·재적재 뒤의 값이 진짜 입력이다** — xlsx 왕복에서 정밀도가 준다
    # (2109901986.7499976 → 2109901986.749998). 원래 리터럴과 대면 헛맞는다.
    back = openpyxl.load_workbook(src, data_only=True).active
    return back, drawn


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--cases", type=int, default=1500, help="반올림 축에서 뽑을 무작위 셀 수")
    ap.add_argument("--seed", type=int, default=31)
    ap.add_argument("--show", type=int, default=10, help="어긋난 것을 몇 개까지 보일지")
    a = ap.parse_args()

    soffice = find_soffice()
    if not soffice:
        # **「못 쟀다」와 「이상 없다」를 같은 종료코드로 두지 않는다.** 조용함의 뜻이
        # 둘이 되면 다음 사람이 안 돌아간 검사를 통과로 읽는다.
        print("✗ LibreOffice(soffice)를 찾지 못해 **재지 못했습니다.** 통과가 아닙니다.")
        print("  찾아본 자리: " + " · ".join(_SOFFICE_CANDIDATES))
        return 2

    cases = build_cases(a.cases, a.seed)
    print(f"LibreOffice: {soffice}")
    print(f"셀 {len(cases)}개 (구역 축 {len(_SECTION_FORMATS)}서식 × {len(_SECTION_VALUES)}값 전수 + 무작위 {a.cases})")

    with tempfile.TemporaryDirectory() as td:
        back, drawn = render_with_libreoffice(soffice, cases, Path(td))
        if len(drawn) != back.max_row:
            print(f"✗ 행 수가 안 맞습니다: csv {len(drawn)} · xlsx {back.max_row}")
            return 1

        same = spaced = unknown = 0
        bad = []
        for i in range(1, back.max_row + 1):
            cell = back.cell(row=i, column=1)
            text, known = render_value(cell)
            want = drawn[i - 1]
            if not known:
                unknown += 1
            elif text == want:
                same += 1
            elif re.sub(r"\s+", "", text) == re.sub(r"\s+", "", want):
                # 회계 서식의 `_-` 는 「뒤 글자 폭만큼 빈칸」이라 LibreOffice 는 공백을
                # 그리고 우리는 뺀다. 표 안에서는 뜻이 없는 정렬용이라 어긋남이 아니다.
                spaced += 1
            else:
                bad.append((repr(cell.value), cell.number_format, text, want))

    tot = back.max_row
    print()
    print(f"  글자까지 일치        {same:>6}건 ({same / tot * 100:5.2f}%)")
    print(f"  정렬 공백만 다름     {spaced:>6}건 ({spaced / tot * 100:5.2f}%)  ← 회계 서식의 `_-`")
    print(f"  안 그리고 셈         {unknown:>6}건 ({unknown / tot * 100:5.2f}%)  ← 갈리는 값. 고장 아님")
    print(f"  숫자가 어긋남        {len(bad):>6}건 ({len(bad) / tot * 100:5.2f}%)  ← 이것만 0 이어야 한다")

    if not bad:
        print("\n✓ 그린 것 중 LibreOffice 와 어긋난 셀이 없습니다.")
        return 0

    print("\n✗ 어긋난 셀:")
    for v, fmt, got, want in bad[: a.show]:
        print(f"    값={v:<26} 서식={fmt:<28} 우리={got!r} LibreOffice={want!r}")
    if len(bad) > a.show:
        print(f"    … 그 외 {len(bad) - a.show}건")
    return 1


if __name__ == "__main__":
    sys.exit(main())
